# -*- coding: utf-8 -*-
"""夏制 · 五行五色 示范: 色卡表 + 环形/柱状图(五行主题) — 照镜自检流程交付"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import Reference
from openpyxl.chart.series import DataPoint
from chart_style import (use_theme, wuxing, styled_donut, styled_bar, set_brand,
                         WUXING, WUXING_N, GraphicalProperties, LineProperties)

use_theme('五行')   # 整套切到五行主题(标题玄青/网格暖灰/宣纸底)
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'output')
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active; ws.title = '五行五色'
ws.sheet_view.showGridLines = False
# 横向单页 + 窄边距(防止绝对锚点被分页裁切)
from openpyxl.worksheet.page import PageMargins
ws.page_setup.orientation = 'landscape'
ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3, header=0.1, footer=0.1)
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.utils.units import cm_to_EMU as C2E
def abs_at(x_cm, y_cm, w_cm, h_cm):
    return AbsoluteAnchor(pos=XDRPoint2D(C2E(x_cm), C2E(y_cm)),
                          ext=XDRPositiveSize2D(C2E(w_cm), C2E(h_cm)))

INK = '37474F'; PAPER = WUXING_N['宣纸底']; CARD = WUXING_N['卡片']
FONT = '微软雅黑'

# ── 标题区 ──
ws['B2'] = '五行五色 · 夏制配色'
ws['B2'].font = Font(size=18, bold=True, color=INK, name=FONT)
ws['B3'] = '中国五行正色的低饱和演绎 —— 耐看、不艳丽，适用于 PPT 与部分表格'
ws['B3'].font = Font(size=10, color=WUXING_N['次字'], name=FONT)

# ── 色卡表 ──
hdr = ['五行', '色名', 'HEX', 'RGB', '用途/气韵']
rows = [
    ('木', '竹青', '#6F8F7C', '111, 143, 124', '主系列 · 生长'),
    ('火', '赭石', '#A2594B', '162, 89, 75',   '强调 / 高亮 · 热烈'),
    ('土', '秋香', '#C7A96A', '199, 169, 106', '次要系列 · 承载'),
    ('金', '月白', '#AEB9BE', '174, 185, 190', '参考 / 目标线 · 收敛'),
    ('水', '玄青', '#37474F', '55, 71, 79',    '标题 / 主色 · 深沉'),
]
r0 = 5
for c, h in enumerate(hdr, start=2):
    cell = ws.cell(row=r0, column=c, value=h)
    cell.font = Font(bold=True, color='FFFFFF', size=10, name=FONT)
    cell.fill = PatternFill('solid', fgColor=INK)
    cell.alignment = Alignment(horizontal='center', vertical='center')
thin = Side(style='thin', color='D8D2C6')
for i, (el, nm, hexv, rgb, use) in enumerate(rows):
    r = r0 + 1 + i
    hx = hexv.lstrip('#')
    # 五行格: 填该色
    c1 = ws.cell(row=r, column=2, value=el)
    c1.fill = PatternFill('solid', fgColor=hx)
    c1.font = Font(bold=True, size=12, color='FFFFFF' if el in ('火', '水') else '3A3A3A', name=FONT)
    c1.alignment = Alignment(horizontal='center', vertical='center')
    c2 = ws.cell(row=r, column=3, value=nm)
    c2.font = Font(size=11, bold=True, color=INK, name=FONT)
    c3 = ws.cell(row=r, column=4, value=hexv)
    c3.font = Font(size=10, color=WUXING_N['次字'], name=FONT)
    c4 = ws.cell(row=r, column=5, value=rgb)
    c4.font = Font(size=10, color=WUXING_N['次字'], name=FONT)
    c5 = ws.cell(row=r, column=6, value=use)
    c5.font = Font(size=10, color=WUXING_N['次字'], name=FONT)
    for c in range(2, 7):
        ws.cell(row=r, column=c).border = Border(bottom=thin)
for col, w in zip('BCDEF', [8, 9, 11, 15, 24]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[r0].height = 22

ws['B12'] = '用法: use_theme("五行") 一键切换整套(标题/网格/轴/底色); 系列色用 wuxing(n)'
ws['B12'].font = Font(size=9, italic=True, color=WUXING_N['次字'], name=FONT)

# ── 数据源(隐藏) ──
wd = wb.create_sheet('源'); wd.sheet_state = 'hidden'
wd.append(['元素', '关注度', '占比'])
data = [('木', 34, 28), ('火', 26, 22), ('土', 24, 20), ('金', 19, 15), ('水', 18, 15)]
for d in data:
    wd.append(list(d))

# ── 图1: 五行环形(五色) ──
ch1 = styled_donut(
    ws,
    Reference(wd, min_col=3, min_row=1, max_col=3, max_row=6),
    Reference(wd, min_col=1, min_row=2, max_row=6),
    wuxing(5),
    '五行占比 · 五色亮相', 'B15', hole=52, width=12.3, height=8.8)
ch1.anchor = abs_at(0.8, 8.8, 12.3, 8.8)
for i, dp in enumerate(ch1.series[0].data_points or []):
    dp.spPr = GraphicalProperties(solidFill=wuxing(5)[i], ln=LineProperties(solidFill='FAF8F4', w=19050))

# ── 图2: 五行柱状(逐点五色) ──
ch2 = styled_bar(
    ws,
    Reference(wd, min_col=2, min_row=1, max_col=2, max_row=6),
    Reference(wd, min_col=1, min_row=2, max_row=6),
    '五行关注度 · 一色一系', 'B15',
    colors=(wuxing(5)[0],), labels=True, legend=False, gap=60)
ch2.anchor = abs_at(13.5, 8.8, 12.3, 8.8)
ch2.series[0].data_points = [
    DataPoint(idx=i, spPr=GraphicalProperties(solidFill=c, ln=LineProperties(noFill=True)))
    for i, c in enumerate(wuxing(5))]

set_brand(wb, title='五行五色 · 夏制配色', subject='五行五色 · 低饱和演绎 · 照镜自检')
wb.save(os.path.join(OUT, 'wuxing_palette.xlsx'))
print('生成:', os.path.join(OUT, 'wuxing_palette.xlsx'))

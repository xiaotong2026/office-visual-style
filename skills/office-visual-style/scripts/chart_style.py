# -*- coding: utf-8 -*-
"""
Excel图表美化助手 v1.0 (2026-09-23)
用途: openpyxl生成原生可编辑Excel图表的统一美化标准
设计原则: 洞见式标题 / 去图表垃圾 / 直接标注 / 低饱和专业配色 / 浅网格层级
"""
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (Paragraph, ParagraphProperties,
                                   CharacterProperties, Font as XFont)
from openpyxl.worksheet.properties import PageSetupProperties

# ── 商务配色(低饱和) ──────────────────────────────
NAVY   = "1F3864"   # 标题
BLUE   = "2E6FB7"   # 主色
ORANGE = "E8873A"   # 强调色(高亮重点)
LIGHT  = "BDD7EE"   # 次要系列
PIE_RAMP = ["E8873A", "2E6FB7", "5B8FD0", "88AFDF", "B5CDEF"]  # 首位=高亮
TXT_DARK, TXT_GRAY = "262626", "595959"
GRID, AXISLINE = "E9E9E9", "C9C9C9"
FONT = "微软雅黑"   # Windows端显示; Linux渲染自动替换Noto Sans CJK


# ══ 夏制 · 五行五色 (Wuxing Five-Color) — 中国五行正色的低饱和演绎,耐看不艳丽 ══
# 用途: PPT / 部分表格的替代配色主题(不替代默认"盛夏蓝橙")
WUXING = {
    "木·竹青": "6F8F7C",   # 生长  → 主系列
    "火·赭石": "A2594B",   # 热烈  → 强调/高亮
    "土·秋香": "C7A96A",   # 承载  → 次要系列
    "金·月白": "AEB9BE",   # 收敛  → 参考/目标(银)
    "水·玄青": "37474F",   # 深沉  → 标题/主色
}
WUXING_SERIES = ["6F8F7C", "A2594B", "C7A96A", "AEB9BE", "37474F"]  # 木→火→土→金→水
WUXING_N = {"宣纸底": "FAF8F4", "墨字": "2B2B2B", "次字": "6B6B6B",
            "网格": "E7E2D9", "轴": "C9C3B6", "卡片": "F6F3EC", "牌匾": "F6F3EC"}


def wuxing(n=None):
    """五行五色系列色(木→火→土→金→水); n=取前n色, None=全部"""
    return WUXING_SERIES[:n] if n is not None else list(WUXING_SERIES)


# ── 主题引擎: 一键切换整套配色(标题/网格/轴/底) ──
_THEME = {"title": NAVY, "grid": GRID, "axis": AXISLINE, "chart_bg": None, "text": TXT_DARK}


def use_theme(name="商务"):
    """切换配色主题: '商务'(默认·盛夏蓝橙) / '五行'(wuxing)"""
    if str(name).lower() in ("五行", "wuxing"):
        _THEME.update(title=WUXING["水·玄青"], grid=WUXING_N["网格"], axis=WUXING_N["轴"],
                      chart_bg=WUXING_N["宣纸底"], text=WUXING_N["墨字"])
    else:
        _THEME.update(title=NAVY, grid=GRID, axis=AXISLINE, chart_bg=None, text=TXT_DARK)
    return dict(_THEME)


def set_theme(**kw):
    _THEME.update({k: v for k, v in kw.items() if k in _THEME})
    return dict(_THEME)


def _rpr(sz=1000, color=TXT_GRAY, bold=False):
    return CharacterProperties(sz=sz, b=bold, solidFill=color,
                               latin=XFont(typeface=FONT), ea=XFont(typeface=FONT))


def _rich(sz=1000, color=TXT_GRAY, bold=False):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=_rpr(sz, color, bold)))])


def set_title(chart, text, sz=1400):
    """洞见式标题: 说结论,不说'统计'"""
    chart.title = text
    col = _THEME.get("title", NAVY)
    para = chart.title.tx.rich.p[0]
    para.pPr = ParagraphProperties(defRPr=_rpr(sz, col, True))
    if para.r:
        for run in para.r:
            run.rPr = _rpr(sz, col, True)


def style_axes(chart, num_fmt='#,##0'):
    """浅网格线 + 弱化轴线 + 统一字号"""
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill=_THEME.get('grid', GRID), w=9525)))
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.y_axis.majorTickMark = 'none'
    chart.y_axis.delete = False
    chart.y_axis.txPr = _rich(1000)
    chart.y_axis.number_format = num_fmt
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=_THEME.get('axis', AXISLINE), w=9525))
    chart.x_axis.majorTickMark = 'none'
    chart.x_axis.delete = False
    chart.x_axis.txPr = _rich(1000)


def clean_frame(chart, legend_pos='b'):
    """去外边框 + 图例置底 (+主题底色)"""
    _bg = _THEME.get("chart_bg")
    if _bg:
        chart.graphical_properties = GraphicalProperties(
            solidFill=_bg, ln=LineProperties(noFill=True))
    else:
        chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    if chart.legend:
        chart.legend.position = legend_pos
        chart.legend.overlay = False
        chart.legend.txPr = _rich(1000)


def add_labels(series, pos='outEnd', num_fmt='#,##0', sz=900):
    """直接标注数据值,省去看轴(显式关闭多余项,防LibreOffice误渲染)"""
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.showLegendKey = False
    series.dLbls.showCatName = False
    series.dLbls.showSerName = False
    series.dLbls.showPercent = False
    series.dLbls.showBubbleSize = False
    series.dLbls.numFmt = num_fmt
    series.dLbls.dLblPos = pos
    series.dLbls.txPr = _rich(sz, _THEME.get('text', TXT_DARK))


def styled_bar(ws, data_ref, cats_ref, title, anchor,
               colors=(BLUE, ORANGE), gap=90, overlap=-10, labels=True, legend=True,
               horizontal=False, num_fmt='#,##0'):
    ch = BarChart(); ch.type = 'bar' if horizontal else 'col'
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    for i, s in enumerate(ch.series):
        s.graphicalProperties = GraphicalProperties(
            solidFill=colors[i % len(colors)], ln=LineProperties(noFill=True))
    ch.gapWidth = gap
    ch.overlap = overlap
    set_title(ch, title)
    style_axes(ch)
    if labels:
        for s in ch.series:
            add_labels(s)
    if not legend:
        ch.legend = None
    clean_frame(ch)
    ch.width, ch.height = 20, 11
    ws.add_chart(ch, anchor)
    return ch


def styled_line(ws, data_ref, cats_ref, title, anchor,
                color=BLUE, width_pt=2.25, labels=True, legend=False):
    ch = LineChart()
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    for s in ch.series:
        s.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=color, w=int(width_pt * 12700)))
        s.smooth = False
        s.marker = Marker(symbol='circle', size=6,
                          spPr=GraphicalProperties(solidFill=color,
                                                   ln=LineProperties(solidFill='FFFFFF', w=12700)))
    set_title(ch, title)
    style_axes(ch)
    if labels:
        for s in ch.series:
            add_labels(s, pos='t')
    if not legend:
        ch.legend = None
    clean_frame(ch)
    ch.width, ch.height = 20, 11
    ws.add_chart(ch, anchor)
    return ch


def styled_pie(ws, data_ref, cats_ref, title, anchor, colors=PIE_RAMP, highlight_first=True):
    ch = PieChart()
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ser = ch.series[0]
    ser.data_points = [DataPoint(idx=i, spPr=GraphicalProperties(
        solidFill=c, ln=LineProperties(solidFill='FFFFFF', w=19050)))
        for i, c in enumerate(colors)]
    ser.dLbls = DataLabelList()
    ser.dLbls.showPercent = True
    ser.dLbls.showVal = False
    ser.dLbls.showLegendKey = False
    ser.dLbls.showCatName = False
    ser.dLbls.showSerName = False
    ser.dLbls.showBubbleSize = False
    ser.dLbls.dLblPos = 'bestFit'
    ser.dLbls.txPr = _rich(950, TXT_DARK, True)
    set_title(ch, title)
    clean_frame(ch)
    ch.width, ch.height = 14, 11
    ws.add_chart(ch, anchor)
    return ch


def fit_to_page(ws, landscape=True):
    """防图表被分页切割(渲染/打印前必设)"""
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


# ══ v1.1 箱形图(2026-09-23) ══
# Excel2016+原生箱形图是chartEx新格式,openpyxl不支持 → 用"堆叠柱+误差线"经典手法造原生可编辑箱形图
def styled_box(ws, cats_ref, refs, title, anchor, err_refs=(None, None),
               lower_color=LIGHT, upper_color=BLUE, whisker=NAVY, gap=80,
               num_fmt='0.0', labels=False):
    """refs=(Q1基准,箱下段Q1→中位,箱上段中位→Q3) err_refs=(下须宽ref,上须宽ref)
    箱体分色边界即中位数;误差线作上下须"""
    from openpyxl.chart import Series
    from openpyxl.chart.error_bar import ErrorBars
    from openpyxl.chart.data_source import NumDataSource, NumRef
    from openpyxl.chart.series import SeriesLabel

    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'stacked'
    ch.overlap = 100
    ch.gapWidth = gap

    for ref in refs:  # 3段: Q1基准/箱下段/箱上段
        ch.add_data(ref, from_rows=True, titles_from_data=False)

    ch.set_categories(cats_ref)

    def _err(ref, direction):
        src = NumDataSource(numRef=NumRef(f=str(ref)))
        kw = dict(errDir='y', errValType='cust', noEndCap=False,
                  spPr=GraphicalProperties(ln=LineProperties(solidFill=whisker, w=15875)))
        if direction == 'minus':
            kw['errBarType'] = 'minus'; kw['minus'] = src
        else:
            kw['errBarType'] = 'plus'; kw['plus'] = src
        return ErrorBars(**kw)

    # 基准段:完全透明,挂下须
    ch.series[0].tx = SeriesLabel(v='Q1基准')
    ch.series[0].graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    if err_refs[0] is not None:
        ch.series[0].errBars = _err(err_refs[0], 'minus')

    # 箱体两段:浅蓝/主蓝,分色边界=中位数
    ch.series[1].tx = SeriesLabel(v='箱体下半(Q1→中位)')
    ch.series[1].graphicalProperties = GraphicalProperties(
        solidFill=lower_color, ln=LineProperties(solidFill='FFFFFF', w=9525))
    ch.series[2].tx = SeriesLabel(v='箱体上半(中位→Q3)')
    ch.series[2].graphicalProperties = GraphicalProperties(
        solidFill=upper_color, ln=LineProperties(solidFill='FFFFFF', w=9525))
    if err_refs[1] is not None:
        ch.series[2].errBars = _err(err_refs[1], 'plus')

    if labels:
        for s in ch.series[1:]:
            add_labels(s, pos='ctr', num_fmt=num_fmt)

    set_title(ch, title)
    style_axes(ch, num_fmt=num_fmt)
    ch.legend = None
    clean_frame(ch)
    ch.width, ch.height = 22, 12
    ws.add_chart(ch, anchor)
    return ch


# ══ v1.2 复杂场景(2026-09-23晚) ══
def styled_combo(ws, bar_ref, line_ref, cats_ref, title, anchor,
                 bar_color=BLUE, line_color=ORANGE, labels=True, y2_fmt='0.0"%"'):
    """双轴组合图: 柱状(主轴,量) + 折线(次轴,率)。
    二次轴配方: lc.y_axis.axId=200 + 主轴crosses='max'"""
    from openpyxl.chart import LineChart
    ch = BarChart(); ch.type = 'col'
    ch.add_data(bar_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    for s in ch.series:
        s.graphicalProperties = GraphicalProperties(solidFill=bar_color, ln=LineProperties(noFill=True))
    ch.gapWidth = 80
    set_title(ch, title); style_axes(ch)
    if labels:
        for s in ch.series: add_labels(s)

    lc = LineChart()
    lc.add_data(line_ref, titles_from_data=True)
    lc.set_categories(cats_ref)
    for s in lc.series:
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=line_color, w=int(2.25*12700)))
        s.smooth = False
        s.marker = Marker(symbol='circle', size=6,
                          spPr=GraphicalProperties(solidFill=line_color, ln=LineProperties(solidFill='FFFFFF', w=12700)))
    lc.y_axis.axId = 200
    lc.y_axis.majorGridlines = None
    lc.y_axis.number_format = y2_fmt
    lc.y_axis.delete = False
    lc.y_axis.majorTickMark = 'none'
    lc.y_axis.txPr = _rich(1000, line_color)
    lc.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.y_axis.crosses = 'max'
    ch += lc
    if ch.legend:
        ch.legend.position = 'b'; ch.legend.overlay = False; ch.legend.txPr = _rich(1000)
    clean_frame(ch)
    ch.width, ch.height = 22, 11
    ws.add_chart(ch, anchor)
    return ch


def styled_waterfall(ws, cats_ref, base_ref, delta_ref, colors, title, anchor,
                     gap=40, num_fmt='#,##0'):
    """瀑布图(利润桥): 透明基准段 + delta段(DataPoint逐点上色)。
    colors=delta逐点色表(蓝=增/橙=减/深蓝=起止合计)。
    delta系列showVal天然=真实金额(增减额/起止值),标注语义正确"""
    from openpyxl.chart import Series
    from openpyxl.chart.series import DataPoint, SeriesLabel

    ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = gap
    for ref in (base_ref, delta_ref):
        ch.add_data(ref, from_rows=True, titles_from_data=False)
    ch.set_categories(cats_ref)

    ch.series[0].tx = SeriesLabel(v='base')
    ch.series[0].graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))

    ser = ch.series[1]
    ser.tx = SeriesLabel(v='delta')
    ser.graphicalProperties = GraphicalProperties(solidFill=BLUE)
    ser.data_points = [DataPoint(idx=i, spPr=GraphicalProperties(
        solidFill=c, ln=LineProperties(noFill=True))) for i, c in enumerate(colors)]
    ser.dLbls = DataLabelList()
    ser.dLbls.showVal = True
    for f in ('showLegendKey', 'showCatName', 'showSerName', 'showPercent', 'showBubbleSize'):
        setattr(ser.dLbls, f, False)
    ser.dLbls.numFmt = num_fmt
    ser.dLbls.txPr = _rich(950, TXT_DARK, True)

    ch.legend = None
    set_title(ch, title); style_axes(ch, num_fmt=num_fmt)
    clean_frame(ch)
    ch.width, ch.height = 22, 11
    ws.add_chart(ch, anchor)
    return ch


# ══ v1.3 高阶场景(2026-09-23夜) ══
def styled_stacked(ws, data_ref, cats_ref, title, anchor,
                   colors=(BLUE, ORANGE, LIGHT), labels=True, legend=True, gap=60):
    """堆叠柱(构成分析): 每列=一个系列(带表头)。labels标注各段值(ctr)"""
    ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = gap
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    for i, s in enumerate(ch.series):
        s.graphicalProperties = GraphicalProperties(
            solidFill=colors[i % len(colors)], ln=LineProperties(solidFill='FFFFFF', w=9525))
    set_title(ch, title); style_axes(ch)
    if labels:
        for s in ch.series:
            add_labels(s, pos='ctr')
    if not legend:
        ch.legend = None
    elif ch.legend:
        ch.legend.position = 'b'; ch.legend.overlay = False; ch.legend.txPr = _rich(1000)
    clean_frame(ch)
    ch.width, ch.height = 20, 11
    ws.add_chart(ch, anchor)
    return ch


# ══ v1.4 管线场景(2026-09-23深夜) ══
def styled_target(ws, actual_ref, target_ref, cats_ref, title, anchor,
                  actual_color=BLUE, target_color=ORANGE, labels=True,
                  width=20, height=11):
    """实际vs目标(同轴): 实际柱 + 目标虚线。目标线用prstDash='dash',同轴无需次轴"""
    from openpyxl.chart import LineChart
    ch = BarChart(); ch.type = 'col'
    ch.add_data(actual_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    for s in ch.series:
        s.graphicalProperties = GraphicalProperties(solidFill=actual_color, ln=LineProperties(noFill=True))
    ch.gapWidth = 70
    set_title(ch, title); style_axes(ch)
    if labels:
        add_labels(ch.series[0])

    lc = LineChart()
    lc.add_data(target_ref, titles_from_data=True)
    lc.set_categories(cats_ref)
    for s in lc.series:
        s.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=target_color, w=int(2.25 * 12700), prstDash='dash'))
        s.smooth = False
        s.marker = Marker(symbol='none')
    ch += lc
    if ch.legend:
        ch.legend.position = 'b'; ch.legend.overlay = False; ch.legend.txPr = _rich(1000)
    clean_frame(ch)
    ch.width, ch.height = width, height
    ws.add_chart(ch, anchor)
    return ch


def styled_donut(ws, data_ref, cats_ref, colors, title, anchor, hole=55,
                 num_fmt=None, width=13, height=11):
    """环形占比图: 传占比数值(如42.1),标签显示百分比"""
    from openpyxl.chart import DoughnutChart
    from openpyxl.chart.series import DataPoint
    ch = DoughnutChart(); ch.holeSize = hole
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ser = ch.series[0]
    ser.data_points = [DataPoint(idx=i, spPr=GraphicalProperties(
        solidFill=c, ln=LineProperties(solidFill='FFFFFF', w=19050)))
        for i, c in enumerate(colors)]
    ser.dLbls = DataLabelList()
    ser.dLbls.showVal = True
    for f in ('showLegendKey', 'showCatName', 'showSerName', 'showPercent', 'showBubbleSize'):
        setattr(ser.dLbls, f, False)
    if num_fmt:
        ser.dLbls.numFmt = num_fmt
    ser.dLbls.txPr = _rich(1000, TXT_DARK, True)
    set_title(ch, title)
    clean_frame(ch)
    ch.width, ch.height = width, height
    ws.add_chart(ch, anchor)
    return ch


def styled_scatter(ws, x_ref, y_ref, title, anchor, color=BLUE,
                   x_title='X', y_title='Y', width=20, height=11):
    """散点图(相关关系): 记得series.line.noFill否则连线"""
    from openpyxl.chart import Series, ScatterChart
    ch = ScatterChart()
    ch.scatterStyle = 'marker'
    ser = Series(y_ref, x_ref, title='观测点')
    ser.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ser.marker = Marker(symbol='circle', size=5,
                        spPr=GraphicalProperties(solidFill=color, ln=LineProperties(noFill=True)))
    ch.series.append(ser)
    set_title(ch, title)
    ch.x_axis.title = x_title
    ch.y_axis.title = y_title
    style_axes(ch)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend = None
    clean_frame(ch)
    ch.width, ch.height = width, height
    ws.add_chart(ch, anchor)
    return ch


# ══════════════════════════════════════════════════════════════
#  夏制 (Xiazhi) — 让 AI 出的图表，像人做的
#  配色系统: 盛夏蓝橙  |  自检闭环: 照镜自检
#  github.com/xiaotong2026/office-visual-style
# ══════════════════════════════════════════════════════════════
BRAND = '夏制'
BRAND_EN = 'Xiazhi'
BRAND_URL = 'github.com/xiaotong2026/office-visual-style'
PALETTE_NAME = '盛夏蓝橙'
LOOP_NAME = '照镜自检'
BRAND_GREY = '9AA5B1'


def set_brand(wb, title=None, subject=None, footer=True):
    """工作簿品牌层: 写入文件属性 + 每页页脚克制署名(作品自带出处,可溯源)"""
    p = wb.properties
    p.creator = f'{BRAND} {BRAND_EN}'
    p.lastModifiedBy = f'{BRAND} {BRAND_EN}'
    p.category = f'{BRAND} · office-visual-style'
    p.keywords = f'{BRAND},{BRAND_EN},{BRAND_URL}'
    p.description = f'Made with {BRAND} ({BRAND_EN}) · {BRAND_URL}'
    if title:
        p.title = title
    if subject:
        p.subject = subject
    if footer:
        for ws in wb.worksheets:
            try:
                ws.oddFooter.right.text = f'{BRAND} · {BRAND_EN}'
                ws.oddFooter.right.size = 8
                ws.oddFooter.right.color = BRAND_GREY
            except Exception:
                pass
    return wb


def brand_ppt(prs, footer=True, text=None):
    """PPT品牌层: 核心属性 + 每页右下角克制署名"""
    cp = prs.core_properties
    cp.author = f'{BRAND} {BRAND_EN}'
    cp.last_modified_by = f'{BRAND} {BRAND_EN}'
    cp.category = f'{BRAND} · office-visual-style'
    cp.keywords = f'{BRAND},{BRAND_EN},office-visual-style'
    cp.comments = f'Made with {BRAND} · {BRAND_URL}'
    if footer:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        for slide in prs.slides:
            tb = slide.shapes.add_textbox(Inches(9.9), Inches(7.05), Inches(3.0), Inches(0.3))
            para = tb.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.RIGHT
            r = para.add_run()
            r.text = text or f'◆ {BRAND} · {BRAND_EN}'
            r.font.size = Pt(9)
            r.font.name = '微软雅黑'
            r.font.color.rgb = RGBColor(0x9A, 0xA5, 0xB1)
    return prs

def wuxing_rgb():
    """返回pptx用五行RGBColor列表(木→火→土→金→水)"""
    from pptx.dml.color import RGBColor
    return [RGBColor.from_string(c) for c in WUXING_SERIES]


def wuxing_ppt_theme(prs, footer_text='◆ 夏制 · 五行五色'):
    """把整套PPT刷成五行主题: 全部slide宣纸底 + 右下品牌署名。系列色用 wuxing_rgb()"""
    from pptx.dml.color import RGBColor
    paper = RGBColor.from_string(WUXING_N['宣纸底'])
    for slide in prs.slides:
        try:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = paper
        except Exception:
            pass
    return brand_ppt(prs, text=footer_text)


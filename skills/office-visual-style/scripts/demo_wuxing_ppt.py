# -*- coding: utf-8 -*-
"""夏制 · 五行五色 PPT样张 (3页) — 上大屏效果演示"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from chart_style import brand_ppt

# 五行五色
WOOD = RGBColor(0x6F, 0x8F, 0x7C); FIRE = RGBColor(0xA2, 0x59, 0x4B)
EARTH = RGBColor(0xC7, 0xA9, 0x6A); METAL = RGBColor(0xAE, 0xB9, 0xBE)
WATER = RGBColor(0x37, 0x47, 0x4F)
PAPER = RGBColor(0xFA, 0xF8, 0xF4); INK = RGBColor(0x2B, 0x2B, 0x2B)
GREY = RGBColor(0x6B, 0x6B, 0x6B); CARD = RGBColor(0xF6, 0xF3, 0xEC)
GRID = RGBColor(0xE7, 0xE2, 0xD9); AXIS = RGBColor(0xC9, 0xC3, 0xB6)
FIVE = [WOOD, FIRE, EARTH, METAL, WATER]
FONT = '微软雅黑'

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'output')
os.makedirs(OUT, exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def txt(s, x, y, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, ls=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb


def paper(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER


# ═ 页1 封面 ═
s = prs.slides.add_slide(BLANK); paper(s)
# 左侧竖色条(五行)
for i, c in enumerate(FIVE):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.2 + i * 0.62), Inches(0.16), Inches(0.52))
    bar.fill.solid(); bar.fill.fore_color.rgb = c; bar.line.fill.background()
txt(s, 1.35, 1.55, 8, 0.4, '夏制 · XIAZHI', 14, FIRE, True)
txt(s, 1.35, 2.15, 10, 1.2, '五行五色', 54, WATER, True)
txt(s, 1.35, 3.45, 10, 0.6, '木火土金水 · 中国正色的低饱和演绎', 20, GREY)
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.35), Inches(4.35), Inches(1.8), Pt(3))
line.fill.solid(); line.fill.fore_color.rgb = FIRE; line.line.fill.background()
txt(s, 1.35, 4.6, 11, 1.2,
    '耐看 · 不艳丽 · 适用 PPT 与部分表格\n配色样张 · 2026', 14, GREY, ls=1.4)

# ═ 页2 柱状(一节一色) ═
s = prs.slides.add_slide(BLANK); paper(s)
txt(s, 0.7, 0.45, 12, 0.7, '五行关注度 · 一色一系', 26, WATER, True)
acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.05), Inches(9.2), Pt(2))
acc.fill.solid(); acc.fill.fore_color.rgb = FIRE; acc.line.fill.background()
cd = CategoryChartData(); cd.categories = ['木', '火', '土', '金', '水']
cd.add_series('关注度', (34, 26, 24, 19, 18))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.35), Inches(8.4), Inches(5.4), cd)
ch = gf.chart
ch.has_legend = False; ch.font.size = Pt(13); ch.font.name = FONT; ch.font.color.rgb = INK
ser = ch.plots[0].series[0]
for i, c in enumerate(FIVE):
    pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = c
    pt.format.line.fill.background()
ch.plots[0].gap_width = 60
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels; dl.number_format = '0'; dl.number_format_is_linked = False
dl.font.size = Pt(12); dl.font.bold = True; dl.font.color.rgb = INK; dl.font.name = FONT
ca = ch.category_axis; ca.major_tick_mark = XL_TICK_MARK.NONE
ca.format.line.color.rgb = AXIS
va = ch.value_axis; va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = GRID
va.major_gridlines.format.line.width = Pt(0.5)
va.format.line.fill.background(); va.major_tick_mark = XL_TICK_MARK.NONE
# 右侧色卡
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(1.35), Inches(3.1), Inches(5.4))
card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = GRID; card.line.width = Pt(0.75)
card.adjustments[0] = 0.05
txt(s, 9.75, 1.6, 2.7, 0.4, '五行五色', 15, WATER, True)
els = [('木 · 竹青', '#6F8F7C', WOOD), ('火 · 赭石', '#A2594B', FIRE),
       ('土 · 秋香', '#C7A96A', EARTH), ('金 · 月白', '#AEB9BE', METAL),
       ('水 · 玄青', '#37474F', WATER)]
for i, (nm, hx, c) in enumerate(els):
    y = 2.15 + i * 0.82
    sw = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.75), Inches(y), Inches(0.5), Inches(0.5))
    sw.fill.solid(); sw.fill.fore_color.rgb = c; sw.line.fill.background(); sw.adjustments[0] = 0.18
    txt(s, 10.4, y + 0.02, 2.0, 0.3, nm, 12.5, INK, True)
    txt(s, 10.4, y + 0.28, 2.0, 0.25, hx, 10, GREY)

# ═ 页3 环形 + 说明 ═
s = prs.slides.add_slide(BLANK); paper(s)
txt(s, 0.7, 0.45, 12, 0.7, '五行占比 · 五色亮相', 26, WATER, True)
acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.05), Inches(9.2), Pt(2))
acc.fill.solid(); acc.fill.fore_color.rgb = FIRE; acc.line.fill.background()
cd = CategoryChartData(); cd.categories = ['木', '火', '土', '金', '水']
cd.add_series('占比', (28, 22, 20, 15, 15))
gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.7), Inches(1.35), Inches(6.8), Inches(5.4), cd)
ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.RIGHT
ch.font.size = Pt(13); ch.font.name = FONT; ch.font.color.rgb = INK
ser = ch.plots[0].series[0]
for i, c in enumerate(FIVE):
    pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = c
    pt.format.line.color.rgb = PAPER; pt.format.line.width = Pt(1.5)
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels; dl.number_format = '0"%"'; dl.number_format_is_linked = False
dl.font.size = Pt(12); dl.font.bold = True; dl.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); dl.font.name = FONT
txt(s, 7.9, 1.6, 4.7, 5.0,
    '用色建议\n\n· 主体系 → 木·竹青（沉稳耐看）\n\n· 强调/高亮 → 火·赭石（唯一暖重色，\n  慎用、点在关键处）\n\n· 次要系列 → 土·秋香\n\n· 参考/目标线 → 金·月白\n\n· 标题与主色 → 水·玄青\n\n底色用宣纸色，网格用暖灰，\n整体近水墨，大屏久看不累。',
    14, INK, ls=1.28)

brand_ppt(prs, text='◆ 夏制 · 五行五色')
prs.save(os.path.join(OUT, 'wuxing_ppt.pptx'))
print('生成:', os.path.join(OUT, 'wuxing_ppt.pptx'))
